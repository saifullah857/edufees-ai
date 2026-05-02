# excel_handler.py - All Excel Operations (CRUD + Formatting)

import os
import uuid
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from config import *


# ═══════════════════════════════════════════════════════════════════════════════
#  HELPERS
# ═══════════════════════════════════════════════════════════════════════════════

def _gen_id(prefix: str) -> str:
    return f"{prefix}-{str(uuid.uuid4())[:8].upper()}"


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _thin_border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)


def _header_font(color="FFFFFF"):
    return Font(bold=True, color=color, size=11)


def _fill(hex_color: str):
    return PatternFill("solid", fgColor=hex_color)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _format_sheet(ws, header_color: str, col_widths: dict):
    """Apply header formatting + column widths to a worksheet."""
    border = _thin_border()
    for row_idx, row in enumerate(ws.iter_rows(), start=1):
        for cell in row:
            cell.border = border
            if row_idx == 1:
                cell.fill      = _fill(header_color)
                cell.font      = _header_font()
                cell.alignment = _center()
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row_idx % 2 == 0:
                    cell.fill = _fill(COLOR_ROW_ALT)

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.row_dimensions[1].height = 22


def _color_status_cells(ws, status_cols: list):
    """Color Approved/Pending cells in given column letters."""
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column_letter in status_cols:
                if cell.value == STATUS_APPROVED:
                    cell.fill = _fill(COLOR_APPROVED)
                    cell.font = Font(bold=True, color="276221")
                elif cell.value == STATUS_PENDING:
                    cell.fill = _fill(COLOR_PENDING)
                    cell.font = Font(bold=True, color="9C5700")


# ═══════════════════════════════════════════════════════════════════════════════
#  COLUMN DEFINITIONS
# ═══════════════════════════════════════════════════════════════════════════════

COLLEGE_COLS = [
    "Student_ID", "Name", "Father_Name", "Program", "Section", "Roll_No",
    "Academic_Year", "Total_Annual_Fee",
    "Install1_Amount", "Install1_Status", "Install1_Date",
    "Install2_Amount", "Install2_Status", "Install2_Date",
    "Install3_Amount", "Install3_Status", "Install3_Date",
    "Install4_Amount", "Install4_Status", "Install4_Date",
    "All_Paid", "Next_Year_Unlocked", "Created_At",
]

UNI_COLS = [
    "Student_ID", "Name", "Father_Name", "Department",
    "Roll_No", "Current_Semester", "Semester_Fee", "Created_At",
]

UNI_FEE_COLS = [
    "Record_ID", "Student_ID", "Student_Name", "Department",
    "Semester", "Semester_Fee",
    "Install1_Amount", "Install1_Status", "Install1_Date",
    "Install2_Amount", "Install2_Status", "Install2_Date",
    "Sem_Complete", "Created_At",
]

COLLEGE_COL_WIDTHS = {
    "A": 16, "B": 18, "C": 18, "D": 16, "E": 10, "F": 12,
    "G": 14, "H": 16,
    "I": 14, "J": 12, "K": 14,
    "L": 14, "M": 12, "N": 14,
    "O": 14, "P": 12, "Q": 14,
    "R": 14, "S": 12, "T": 14,
    "U": 10, "V": 16, "W": 18,
}

UNI_COL_WIDTHS = {
    "A": 16, "B": 18, "C": 18, "D": 20,
    "E": 14, "F": 14, "G": 16, "H": 18,
}

UNI_FEE_COL_WIDTHS = {
    "A": 16, "B": 16, "C": 18, "D": 20,
    "E": 10, "F": 14,
    "G": 14, "H": 12, "I": 14,
    "J": 14, "K": 12, "L": 14,
    "M": 14, "N": 18,
}


# ═══════════════════════════════════════════════════════════════════════════════
#  INITIALISE EXCEL
# ═══════════════════════════════════════════════════════════════════════════════

def init_excel():
    """Create Excel file with all sheets if it doesn't already exist."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()

        ws_c = wb.active
        ws_c.title = SHEET_COLLEGE
        ws_c.append(COLLEGE_COLS)
        _format_sheet(ws_c, COLOR_HEADER_COLLEGE, COLLEGE_COL_WIDTHS)

        ws_u = wb.create_sheet(SHEET_UNI)
        ws_u.append(UNI_COLS)
        _format_sheet(ws_u, COLOR_HEADER_UNI, UNI_COL_WIDTHS)

        ws_f = wb.create_sheet(SHEET_UNI_FEE)
        ws_f.append(UNI_FEE_COLS)
        _format_sheet(ws_f, COLOR_HEADER_FEE, UNI_FEE_COL_WIDTHS)

        wb.save(EXCEL_FILE)


def refresh_formatting():
    """Re-apply all cell formatting (call after bulk updates)."""
    wb = load_workbook(EXCEL_FILE)

    ws_c = wb[SHEET_COLLEGE]
    _format_sheet(ws_c, COLOR_HEADER_COLLEGE, COLLEGE_COL_WIDTHS)
    _color_status_cells(ws_c, ["J", "M", "P", "S"])

    ws_u = wb[SHEET_UNI]
    _format_sheet(ws_u, COLOR_HEADER_UNI, UNI_COL_WIDTHS)

    ws_f = wb[SHEET_UNI_FEE]
    _format_sheet(ws_f, COLOR_HEADER_FEE, UNI_FEE_COL_WIDTHS)
    _color_status_cells(ws_f, ["H", "K"])

    for row in ws_f.iter_rows(min_row=2):
        cell = row[12]   # column M → Sem_Complete
        if cell.value == "Yes":
            cell.fill = _fill(COLOR_COMPLETE)
            cell.font = Font(bold=True, color="1F4E79")

    wb.save(EXCEL_FILE)


# ═══════════════════════════════════════════════════════════════════════════════
#  COLLEGE STUDENTS
# ═══════════════════════════════════════════════════════════════════════════════

def add_college_student(
    name: str,
    father_name: str,
    program: str,          # FA / ICS / ICOM / FSC Medical / FSC Engineering / IT
    section: str,
    roll_no: str,
    academic_year: str,
    total_annual_fee: float,
) -> dict:
    """Add a new college student with 4 pending installments."""
    wb  = load_workbook(EXCEL_FILE)
    ws  = wb[SHEET_COLLEGE]
    sid = _gen_id("COL")
    amt = round(total_annual_fee / 4, 2)

    ws.append([
        sid, name, father_name, program, section, roll_no,
        academic_year, total_annual_fee,
        amt, STATUS_PENDING, "",
        amt, STATUS_PENDING, "",
        amt, STATUS_PENDING, "",
        amt, STATUS_PENDING, "",
        "No", "No", _now(),
    ])
    wb.save(EXCEL_FILE)
    refresh_formatting()
    return {
        "success": True,
        "student_id": sid,
        "message": f"College student '{name}' ({program}) added successfully.",
    }


def get_college_students(program_filter=None, section_filter=None) -> pd.DataFrame:
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_COLLEGE, dtype=str)
        # Support legacy "Class" column if present
        if "Class" in df.columns and "Program" not in df.columns:
            df.rename(columns={"Class": "Program"}, inplace=True)
        if program_filter:
            df = df[df["Program"] == program_filter]
        if section_filter:
            df = df[df["Section"] == section_filter]
        return df.fillna("")
    except Exception:
        return pd.DataFrame(columns=COLLEGE_COLS)


def update_college_student(student_id: str, **kwargs) -> dict:
    wb      = load_workbook(EXCEL_FILE)
    ws      = wb[SHEET_COLLEGE]
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2):
        if row[0].value == student_id:
            for key, val in kwargs.items():
                if key in headers:
                    row[headers.index(key)].value = val
            wb.save(EXCEL_FILE)
            refresh_formatting()
            return {"success": True, "message": f"Student {student_id} updated."}
    return {"success": False, "message": "Student not found."}


def delete_college_student(student_id: str) -> dict:
    wb = load_workbook(EXCEL_FILE)
    ws = wb[SHEET_COLLEGE]

    for row in ws.iter_rows(min_row=2):
        if row[0].value == student_id:
            ws.delete_rows(row[0].row)
            wb.save(EXCEL_FILE)
            refresh_formatting()
            return {"success": True, "message": f"Student {student_id} deleted."}
    return {"success": False, "message": "Student not found."}


def approve_college_installment(student_id: str, install_no: int) -> dict:
    """Approve one installment (1–4) for a college student."""
    if install_no not in [1, 2, 3, 4]:
        return {"success": False, "message": "Invalid installment number. Must be 1–4."}

    wb      = load_workbook(EXCEL_FILE)
    ws      = wb[SHEET_COLLEGE]
    headers = [cell.value for cell in ws[1]]

    status_col = f"Install{install_no}_Status"
    date_col   = f"Install{install_no}_Date"

    for row in ws.iter_rows(min_row=2):
        if row[0].value == student_id:
            s_idx = headers.index(status_col)
            d_idx = headers.index(date_col)

            if row[s_idx].value == STATUS_APPROVED:
                return {"success": False, "message": f"Installment {install_no} already approved."}

            row[s_idx].value = STATUS_APPROVED
            row[d_idx].value = _now()

            all_paid = all(
                row[headers.index(f"Install{i}_Status")].value == STATUS_APPROVED
                for i in range(1, 5)
            )
            if all_paid:
                row[headers.index("All_Paid")].value          = "Yes"
                row[headers.index("Next_Year_Unlocked")].value = "Yes"

            wb.save(EXCEL_FILE)
            refresh_formatting()

            msg = f"Installment {install_no} approved for {student_id}."
            if all_paid:
                msg += " ✅ All installments paid — next year unlocked!"
            return {"success": True, "message": msg, "all_paid": all_paid}

    return {"success": False, "message": "Student not found."}


# ═══════════════════════════════════════════════════════════════════════════════
#  UNIVERSITY STUDENTS
# ═══════════════════════════════════════════════════════════════════════════════

def add_university_student(
    name: str,
    father_name: str,
    department: str,
    roll_no: str,
    semester_fee: float,
    starting_semester: int = 1,       # ← NEW: admin can choose starting semester
) -> dict:
    """
    Add a university student and create a fee record for their starting semester.
    semester_fee is looked up from DEFAULT_UNI_FEES if 0 is passed.
    """
    # Auto-fill fee from defaults if not provided
    if semester_fee <= 0:
        semester_fee = DEFAULT_UNI_FEES.get(department, 10_000)

    starting_semester = max(1, min(int(starting_semester), MAX_SEMESTERS))

    wb         = load_workbook(EXCEL_FILE)
    student_id = _gen_id("UNI")

    # University_Students sheet
    ws_u = wb[SHEET_UNI]
    ws_u.append([
        student_id, name, father_name, department,
        roll_no, starting_semester, semester_fee, _now(),
    ])

    # First fee record for chosen starting semester
    ws_f      = wb[SHEET_UNI_FEE]
    record_id = _gen_id("UFR")
    inst_amt  = round(semester_fee / 2, 2)
    ws_f.append([
        record_id, student_id, name, department,
        starting_semester, semester_fee,
        inst_amt, STATUS_PENDING, "",
        inst_amt, STATUS_PENDING, "",
        "No", _now(),
    ])

    wb.save(EXCEL_FILE)
    refresh_formatting()
    return {
        "success": True,
        "student_id": student_id,
        "message": (
            f"University student '{name}' ({department}) added. "
            f"Enrolled from Semester {starting_semester} — "
            f"fee record created (Rs. {semester_fee:,.0f}/sem)."
        ),
    }


def get_university_students(department_filter=None) -> pd.DataFrame:
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_UNI, dtype=str)
        if department_filter:
            df = df[df["Department"] == department_filter]
        return df.fillna("")
    except Exception:
        return pd.DataFrame(columns=UNI_COLS)


def get_university_fee_records(
    student_id: str  = None,
    name: str        = None,
    department: str  = None,
    semester: int    = None,
) -> pd.DataFrame:
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_UNI_FEE, dtype=str)
        if student_id:
            df = df[df["Student_ID"] == student_id]
        if name:
            df = df[df["Student_Name"].str.lower().str.contains(name.lower(), na=False)]
        if department:
            df = df[df["Department"] == department]
        if semester:
            df = df[df["Semester"] == str(semester)]
        return df.fillna("")
    except Exception:
        return pd.DataFrame(columns=UNI_FEE_COLS)


def update_university_student(student_id: str, **kwargs) -> dict:
    wb      = load_workbook(EXCEL_FILE)
    ws      = wb[SHEET_UNI]
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2):
        if row[0].value == student_id:
            for key, val in kwargs.items():
                if key in headers:
                    row[headers.index(key)].value = val
            wb.save(EXCEL_FILE)
            refresh_formatting()
            return {"success": True, "message": f"Student {student_id} updated."}
    return {"success": False, "message": "Student not found."}


def delete_university_student(student_id: str) -> dict:
    wb    = load_workbook(EXCEL_FILE)
    ws_u  = wb[SHEET_UNI]
    found = False

    for row in ws_u.iter_rows(min_row=2):
        if row[0].value == student_id:
            ws_u.delete_rows(row[0].row)
            found = True
            break

    if not found:
        return {"success": False, "message": "Student not found."}

    ws_f = wb[SHEET_UNI_FEE]
    for r in sorted(
        [row[0].row for row in ws_f.iter_rows(min_row=2) if row[1].value == student_id],
        reverse=True,
    ):
        ws_f.delete_rows(r)

    wb.save(EXCEL_FILE)
    refresh_formatting()
    return {"success": True, "message": f"Student {student_id} and all fee records deleted."}


def approve_uni_installment(record_id: str, install_no: int) -> dict:
    """Approve installment 1 or 2 for a university fee record."""
    if install_no not in [1, 2]:
        return {"success": False, "message": "Invalid installment number. Must be 1 or 2."}

    wb      = load_workbook(EXCEL_FILE)
    ws_f    = wb[SHEET_UNI_FEE]
    headers = [cell.value for cell in ws_f[1]]

    status_col = f"Install{install_no}_Status"
    date_col   = f"Install{install_no}_Date"

    for row in ws_f.iter_rows(min_row=2):
        if row[0].value == record_id:
            s_idx = headers.index(status_col)
            d_idx = headers.index(date_col)

            if row[s_idx].value == STATUS_APPROVED:
                return {"success": False, "message": f"Installment {install_no} already approved."}

            row[s_idx].value = STATUS_APPROVED
            row[d_idx].value = _now()

            both_paid = all(
                row[headers.index(f"Install{i}_Status")].value == STATUS_APPROVED
                for i in range(1, 3)
            )

            sem_advanced = False
            if both_paid:
                row[headers.index("Sem_Complete")].value = "Yes"

                student_id  = row[headers.index("Student_ID")].value
                name        = row[headers.index("Student_Name")].value
                department  = row[headers.index("Department")].value
                current_sem = int(row[headers.index("Semester")].value)
                sem_fee     = float(row[headers.index("Semester_Fee")].value)

                if current_sem < MAX_SEMESTERS:
                    next_sem  = current_sem + 1
                    inst_amt  = round(sem_fee / 2, 2)
                    new_rec   = _gen_id("UFR")
                    ws_f.append([
                        new_rec, student_id, name, department,
                        next_sem, sem_fee,
                        inst_amt, STATUS_PENDING, "",
                        inst_amt, STATUS_PENDING, "",
                        "No", _now(),
                    ])
                    # Advance semester in University_Students
                    ws_u      = wb[SHEET_UNI]
                    u_headers = [c.value for c in ws_u[1]]
                    for u_row in ws_u.iter_rows(min_row=2):
                        if u_row[0].value == student_id:
                            u_row[u_headers.index("Current_Semester")].value = next_sem
                            break
                    sem_advanced = True

            wb.save(EXCEL_FILE)
            refresh_formatting()

            msg = f"Installment {install_no} approved."
            if both_paid and sem_advanced:
                msg += f" ✅ Semester {current_sem} complete — Semester {current_sem + 1} started."
            elif both_paid and not sem_advanced:
                msg += " 🎓 All 8 semesters complete — degree program finished!"
            return {"success": True, "message": msg, "sem_advanced": sem_advanced}

    return {"success": False, "message": "Fee record not found."}


# ═══════════════════════════════════════════════════════════════════════════════
#  DASHBOARD STATS
# ═══════════════════════════════════════════════════════════════════════════════

def get_dashboard_stats() -> dict:
    try:
        col_df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_COLLEGE,  dtype=str)
        uni_df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_UNI,      dtype=str)
        fee_df = pd.read_excel(EXCEL_FILE, sheet_name=SHEET_UNI_FEE,  dtype=str)

        total_college   = len(col_df)
        total_uni       = len(uni_df)
        college_paid    = len(col_df[col_df.get("All_Paid",    pd.Series()) == "Yes"]) if "All_Paid"    in col_df else 0
        uni_sem_done    = len(fee_df[fee_df.get("Sem_Complete", pd.Series()) == "Yes"]) if "Sem_Complete" in fee_df else 0
        pending_records = len(fee_df[fee_df.get("Sem_Complete", pd.Series()) == "No"])  if "Sem_Complete" in fee_df else 0

        # Per-program breakdown for college
        program_col = "Program" if "Program" in col_df.columns else "Class"
        program_breakdown = (
            col_df.groupby(program_col).size().to_dict() if not col_df.empty else {}
        )

        # Per-department breakdown for university
        dept_breakdown = (
            uni_df.groupby("Department").size().to_dict() if not uni_df.empty else {}
        )

        return {
            "total_college":     total_college,
            "total_uni":         total_uni,
            "college_paid":      college_paid,
            "uni_sem_done":      uni_sem_done,
            "pending_records":   pending_records,
            "program_breakdown": program_breakdown,
            "dept_breakdown":    dept_breakdown,
        }
    except Exception:
        return {
            "total_college": 0, "total_uni": 0,
            "college_paid": 0,  "uni_sem_done": 0, "pending_records": 0,
            "program_breakdown": {}, "dept_breakdown": {},
        }